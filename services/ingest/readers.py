"""Hexis ingestion — split from the former services/ingest.py (#89).
Module: readers.
"""

from __future__ import annotations

import argparse
import asyncio
import hashlib
import json
import os
import re
import sys
import time
from dataclasses import dataclass, field
from datetime import datetime, timezone
from enum import Enum
from pathlib import Path
from typing import Any, Callable, Optional
from uuid import UUID

from .config import Anchor, Section, _extract_title

# =========================================================================
# DOCUMENT READERS
# =========================================================================


@dataclass
class ReaderWarning:
    """A structured extraction warning: recorded in source_extraction_runs
    and surfaced wherever the document is opened. Codes in use:
    truncated_rows, ocr_used, ocr_unavailable, image_only_page,
    unsupported_feature, fallback_parser, empty_extraction."""

    code: str
    message: str
    detail: dict[str, Any] = field(default_factory=dict)

    def to_dict(self) -> dict[str, Any]:
        payload: dict[str, Any] = {"code": self.code, "message": self.message}
        if self.detail:
            payload["detail"] = self.detail
        return payload


@dataclass
class ReaderResult:
    """What a reader produces: normalized text plus everything needed for
    durable preservation and citation — structural anchors for locators,
    structured warnings for the extraction run, and (where the reader is the
    only one who sees them, e.g. web fetches) the original artifact bytes."""

    text: str
    warnings: list[ReaderWarning] = field(default_factory=list)
    anchors: list[Anchor] = field(default_factory=list)
    artifact_bytes: bytes | None = None
    artifact_mime: str | None = None
    extractor_name: str = ""
    extractor_version: str = ""
    metadata: dict[str, Any] = field(default_factory=dict)

    def warnings_payload(self) -> list[dict[str, Any]]:
        return [w.to_dict() for w in self.warnings]


def _module_version(module: Any) -> str:
    return str(getattr(module, "__version__", "") or "")


class DocumentReader:
    @staticmethod
    def read(file_path: Path) -> str:
        raise NotImplementedError

    @classmethod
    def read_result(cls, file_path: Path) -> ReaderResult:
        """Rich read: text + warnings + anchors. Default wraps the legacy
        read(); format-aware readers override to add locators and warnings."""
        return ReaderResult(text=cls.read(file_path), extractor_name=cls.__name__)


class MarkdownReader(DocumentReader):
    @staticmethod
    def read(file_path: Path) -> str:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            return f.read()


class TextReader(DocumentReader):
    @staticmethod
    def read(file_path: Path) -> str:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            return f.read()


class CodeReader(DocumentReader):
    LANGUAGE_MAP = {
        ".py": "python",
        ".js": "javascript",
        ".ts": "typescript",
        ".jsx": "javascript-react",
        ".tsx": "typescript-react",
        ".java": "java",
        ".c": "c",
        ".cpp": "cpp",
        ".h": "c-header",
        ".hpp": "cpp-header",
        ".go": "go",
        ".rs": "rust",
        ".rb": "ruby",
        ".php": "php",
        ".swift": "swift",
        ".kt": "kotlin",
        ".scala": "scala",
        ".r": "r",
        ".sql": "sql",
        ".sh": "bash",
        ".bash": "bash",
        ".zsh": "zsh",
        ".ps1": "powershell",
        ".yaml": "yaml",
        ".yml": "yaml",
        ".json": "json",
        ".xml": "xml",
        ".html": "html",
        ".css": "css",
        ".scss": "scss",
        ".less": "less",
    }

    @classmethod
    def read(cls, file_path: Path) -> str:
        language = cls.LANGUAGE_MAP.get(file_path.suffix.lower(), "unknown")
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            content = f.read()
        return f"[Language: {language}]\n[File: {file_path.name}]\n\n{content}"


class WebReader(DocumentReader):
    """Reader for web content via URL."""

    @staticmethod
    def read(url: str) -> str:
        return WebReader.read_result(url).text

    @classmethod
    def read_result(cls, url: str) -> ReaderResult:  # type: ignore[override]
        try:
            import trafilatura
        except ImportError:
            raise RuntimeError(
                "Web reader requires trafilatura: pip install hexis[web]"
            )

        downloaded = trafilatura.fetch_url(url)
        if not downloaded:
            raise RuntimeError(f"Failed to fetch URL: {url}")

        content = trafilatura.extract(downloaded, include_tables=True, include_links=True)
        if not content:
            raise RuntimeError(f"Failed to extract content from URL: {url}")

        # Try to get metadata
        metadata = trafilatura.extract_metadata(downloaded)
        title = getattr(metadata, "title", None) if metadata else None
        author = getattr(metadata, "author", None) if metadata else None
        date = getattr(metadata, "date", None) if metadata else None

        header_parts = [f"[Source: {url}]"]
        if title:
            header_parts.append(f"[Title: {title}]")
        if author:
            header_parts.append(f"[Author: {author}]")
        if date:
            header_parts.append(f"[Date: {date}]")

        fetched_at = datetime.now(timezone.utc).isoformat()
        return ReaderResult(
            text="\n".join(header_parts) + "\n\n" + content,
            # The fetched HTML is the original artifact — the reader is the
            # only place that ever sees it.
            artifact_bytes=downloaded.encode("utf-8", errors="replace")
            if isinstance(downloaded, str) else bytes(downloaded),
            artifact_mime="text/html",
            extractor_name="trafilatura",
            extractor_version=_module_version(trafilatura),
            metadata={
                "url": url,
                "fetched_at": fetched_at,
                "title": title,
                "author": author,
                "date": str(date) if date else None,
            },
        )


class YouTubeTranscriptReader(DocumentReader):
    """Reader for YouTube video transcripts."""

    # Patterns that identify YouTube URLs
    _PATTERNS = [
        re.compile(r"(?:youtube\.com/watch\?.*v=|youtu\.be/)([\w-]{11})"),
    ]

    @classmethod
    def extract_video_id(cls, url: str) -> str | None:
        for pat in cls._PATTERNS:
            m = pat.search(url)
            if m:
                return m.group(1)
        return None

    @classmethod
    def can_handle(cls, url: str) -> bool:
        return cls.extract_video_id(url) is not None

    @staticmethod
    def read(url: str) -> str:  # type: ignore[override]
        video_id = YouTubeTranscriptReader.extract_video_id(url)
        if not video_id:
            raise RuntimeError(f"Could not extract YouTube video ID from: {url}")

        try:
            from youtube_transcript_api import YouTubeTranscriptApi
        except ImportError:
            raise RuntimeError(
                "YouTube transcript reader requires youtube-transcript-api: "
                "pip install youtube-transcript-api"
            )

        try:
            transcript_list = YouTubeTranscriptApi.list_transcripts(video_id)
            # Prefer manually created transcripts, fall back to auto-generated
            try:
                transcript = transcript_list.find_manually_created_transcript(["en"])
            except Exception:
                transcript = transcript_list.find_generated_transcript(["en"])

            entries = transcript.fetch()
            # Build readable transcript
            parts = [f"[Source: {url}]", f"[Format: YouTube Transcript]", f"[Video ID: {video_id}]", ""]
            for entry in entries:
                text = entry.get("text", entry) if isinstance(entry, dict) else str(entry)
                parts.append(str(text))

            return "\n".join(parts)
        except Exception as e:
            raise RuntimeError(f"Failed to fetch YouTube transcript: {e}")


class DataReader(DocumentReader):
    """Reader for structured data files (JSON, YAML, CSV, XML)."""

    DATA_EXTENSIONS = {".json", ".yaml", ".yml", ".csv", ".xml"}

    @classmethod
    def read(cls, file_path: Path) -> str:
        suffix = file_path.suffix.lower()
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            content = f.read()

        format_name = {
            ".json": "JSON",
            ".yaml": "YAML",
            ".yml": "YAML",
            ".csv": "CSV",
            ".xml": "XML",
        }.get(suffix, "Data")

        structure_desc = cls._describe_structure(content, suffix)
        return f"[Format: {format_name}]\n[File: {file_path.name}]\n\n{structure_desc}\n\n--- Raw Content ---\n{content}"

    @classmethod
    def _describe_structure(cls, content: str, suffix: str) -> str:
        """Generate a structural description of the data."""
        try:
            if suffix == ".json":
                import json

                data = json.loads(content)
                return cls._describe_json_structure(data)
            elif suffix in {".yaml", ".yml"}:
                try:
                    import yaml

                    data = yaml.safe_load(content)
                    return cls._describe_json_structure(data)
                except ImportError:
                    return "[YAML parsing unavailable]"
            elif suffix == ".csv":
                return cls._describe_csv_structure(content)
            elif suffix == ".xml":
                return cls._describe_xml_structure(content)
        except Exception as e:
            return f"[Structure analysis failed: {e}]"
        return ""

    @classmethod
    def _describe_json_structure(cls, data: Any, depth: int = 0, max_depth: int = 3) -> str:
        """Describe the structure of JSON/YAML data."""
        indent = "  " * depth
        if depth >= max_depth:
            return f"{indent}..."

        if isinstance(data, dict):
            if not data:
                return f"{indent}{{}}"
            lines = [f"{indent}Object with {len(data)} keys:"]
            for key, value in list(data.items())[:10]:
                value_type = type(value).__name__
                if isinstance(value, dict):
                    lines.append(f"{indent}  - {key}: object ({len(value)} keys)")
                elif isinstance(value, list):
                    lines.append(f"{indent}  - {key}: array ({len(value)} items)")
                else:
                    lines.append(f"{indent}  - {key}: {value_type}")
            if len(data) > 10:
                lines.append(f"{indent}  ... and {len(data) - 10} more keys")
            return "\n".join(lines)
        elif isinstance(data, list):
            if not data:
                return f"{indent}[]"
            lines = [f"{indent}Array with {len(data)} items"]
            if data:
                first = data[0]
                if isinstance(first, dict):
                    lines.append(f"{indent}  Item type: object with keys: {list(first.keys())[:5]}")
                else:
                    lines.append(f"{indent}  Item type: {type(first).__name__}")
            return "\n".join(lines)
        else:
            return f"{indent}{type(data).__name__}: {str(data)[:100]}"

    @classmethod
    def _describe_csv_structure(cls, content: str) -> str:
        """Describe CSV structure."""
        lines = content.strip().split("\n")
        if not lines:
            return "[Empty CSV]"

        header_line = lines[0]
        columns = [col.strip().strip('"').strip("'") for col in header_line.split(",")]

        desc = [f"CSV with {len(columns)} columns and {len(lines) - 1} data rows"]
        desc.append(f"Columns: {', '.join(columns[:10])}")
        if len(columns) > 10:
            desc.append(f"  ... and {len(columns) - 10} more columns")

        return "\n".join(desc)

    @classmethod
    def _describe_xml_structure(cls, content: str) -> str:
        """Describe XML structure (basic)."""
        import re

        root_match = re.search(r"<(\w+)[>\s]", content)
        root_tag = root_match.group(1) if root_match else "unknown"

        tag_counts: dict[str, int] = {}
        for tag in re.findall(r"<(\w+)[>\s]", content):
            tag_counts[tag] = tag_counts.get(tag, 0) + 1

        top_tags = sorted(tag_counts.items(), key=lambda x: -x[1])[:10]

        desc = [f"XML document with root element: <{root_tag}>"]
        desc.append(f"Top elements: {', '.join(f'{t}({c})' for t, c in top_tags)}")

        return "\n".join(desc)


class PDFReader(DocumentReader):
    OCR_DPI = 200

    @staticmethod
    def read(file_path: Path) -> str:
        return PDFReader.read_result(file_path).text

    @classmethod
    def read_result(cls, file_path: Path) -> ReaderResult:
        try:
            import pdfplumber
        except ImportError:
            raise RuntimeError(
                "PDF reader requires pdfplumber: pip install hexis[pdf]"
            )

        warnings: list[ReaderWarning] = []
        anchors: list[Anchor] = []
        parts: list[str] = []
        ocr = cls._ocr_backend(warnings)
        ocr_pages: list[int] = []
        image_only_pages: list[int] = []
        pos = 0

        with pdfplumber.open(file_path) as pdf:
            for i, page in enumerate(pdf.pages):
                page_no = i + 1
                page_text = (page.extract_text() or "").strip()
                if not page_text and ocr is not None:
                    page_text = cls._ocr_page(file_path, i, ocr, warnings)
                    if page_text:
                        ocr_pages.append(page_no)
                if not page_text:
                    image_only_pages.append(page_no)
                    page_text = "[No extractable text on this page]"
                piece = f"[Page {page_no}]\n{page_text}"
                if parts:
                    pos += 2  # "\n\n" separator
                anchors.append(Anchor(kind="page", value=page_no, char_offset=pos))
                parts.append(piece)
                pos += len(piece)

        if ocr_pages:
            warnings.append(ReaderWarning(
                code="ocr_used",
                message=f"OCR extracted text from {len(ocr_pages)} page(s) with no text layer",
                detail={"pages": ocr_pages},
            ))
        if image_only_pages:
            warnings.append(ReaderWarning(
                code="image_only_page",
                message=f"{len(image_only_pages)} page(s) have no extractable text",
                detail={"pages": image_only_pages},
            ))

        return ReaderResult(
            text="\n\n".join(parts),
            warnings=warnings,
            anchors=anchors,
            extractor_name="pdfplumber",
            extractor_version=_module_version(pdfplumber),
        )

    @classmethod
    def _ocr_backend(cls, warnings: list[ReaderWarning]):
        """OCR needs pypdfium2 (rasterize, no system deps) + pytesseract
        (tesseract binary). Missing pieces degrade to a warning, never a
        crash — the pages are flagged image_only instead."""
        try:
            import pypdfium2  # noqa: F401
            import pytesseract  # noqa: F401
        except ImportError:
            warnings.append(ReaderWarning(
                code="ocr_unavailable",
                message="OCR fallback unavailable (pip install hexis[ocr] and the tesseract binary)",
            ))
            return None
        return True

    @classmethod
    def _ocr_page(cls, file_path: Path, page_index: int, ocr, warnings: list[ReaderWarning]) -> str:
        import pypdfium2 as pdfium
        import pytesseract

        try:
            doc = pdfium.PdfDocument(str(file_path))
            try:
                bitmap = doc[page_index].render(scale=cls.OCR_DPI / 72)
                image = bitmap.to_pil()
            finally:
                doc.close()
            return (pytesseract.image_to_string(image) or "").strip()
        except Exception as exc:
            if not any(w.code == "ocr_unavailable" for w in warnings):
                warnings.append(ReaderWarning(
                    code="ocr_unavailable",
                    message=f"OCR failed: {exc}",
                    detail={"page": page_index + 1},
                ))
            return ""


class ImageReader(DocumentReader):
    """Reader for images using OCR."""

    IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tiff", ".webp"}

    @classmethod
    def read(cls, file_path: Path) -> str:
        return cls.read_result(file_path).text

    @classmethod
    def read_result(cls, file_path: Path) -> ReaderResult:
        try:
            import pytesseract
            from PIL import Image
        except ImportError:
            raise RuntimeError(
                "Image reader requires Pillow and pytesseract: pip install hexis[ocr] "
                "(plus the tesseract binary)"
            )

        warnings: list[ReaderWarning] = []
        try:
            image = Image.open(file_path)
            text = pytesseract.image_to_string(image)
        except Exception as e:
            raise RuntimeError(f"OCR failed for {file_path.name}: {e}")
        if not text.strip():
            warnings.append(ReaderWarning(
                code="empty_extraction",
                message="No text detected via OCR",
            ))
            body = f"[Image: {file_path.name}]\n[No text detected via OCR]"
        else:
            warnings.append(ReaderWarning(
                code="ocr_used",
                message="Text extracted via OCR",
            ))
            body = f"[Image: {file_path.name}]\n[OCR Extracted Text]\n\n{text}"
        return ReaderResult(
            text=body,
            warnings=warnings,
            extractor_name="pytesseract",
            extractor_version=_module_version(pytesseract),
        )


class AudioReader(DocumentReader):
    """Reader for audio files using speech-to-text."""

    AUDIO_EXTENSIONS = {".mp3", ".wav", ".m4a", ".ogg", ".flac", ".aac", ".wma"}

    @classmethod
    def read(cls, file_path: Path) -> str:
        try:
            import whisper
        except ImportError:
            raise RuntimeError(
                "Audio reader requires openai-whisper: pip install hexis[media]"
            )

        try:
            model = whisper.load_model("base")
            result = model.transcribe(str(file_path))
            text = result.get("text", "")
            if not text.strip():
                return f"[Audio: {file_path.name}]\n[No speech detected]"
            return f"[Audio: {file_path.name}]\n[Transcription]\n\n{text}"
        except Exception as e:
            return f"[Audio: {file_path.name}]\n[Transcription failed: {e}]"


class VideoReader(DocumentReader):
    """Reader for video files - extracts audio and transcribes."""

    VIDEO_EXTENSIONS = {".mp4", ".avi", ".mov", ".mkv", ".webm", ".wmv", ".flv"}

    @classmethod
    def read(cls, file_path: Path) -> str:
        try:
            from moviepy.editor import VideoFileClip
        except ImportError:
            raise RuntimeError(
                "Video reader requires moviepy: pip install hexis[media]"
            )

        try:
            import whisper
        except ImportError:
            raise RuntimeError(
                "Video reader requires openai-whisper: pip install hexis[media]"
            )

        import tempfile

        try:
            # Extract audio from video
            video = VideoFileClip(str(file_path))
            duration = video.duration

            with tempfile.NamedTemporaryFile(suffix=".wav", delete=False) as tmp_audio:
                audio_path = tmp_audio.name
                video.audio.write_audiofile(audio_path, verbose=False, logger=None)
                video.close()

            # Transcribe the audio
            model = whisper.load_model("base")
            result = model.transcribe(audio_path)
            text = result.get("text", "")

            # Clean up temp file
            import os
            os.unlink(audio_path)

            if not text.strip():
                return f"[Video: {file_path.name}]\n[Duration: {duration:.1f}s]\n[No speech detected]"

            return f"[Video: {file_path.name}]\n[Duration: {duration:.1f}s]\n[Transcription]\n\n{text}"

        except Exception as e:
            return f"[Video: {file_path.name}]\n[Transcription failed: {e}]"


class DocxReader(DocumentReader):
    """Reader for Microsoft Word .docx files."""

    EXTENSIONS = {".docx"}

    _HEADING_RE = re.compile(r"^Heading (\d+)$")

    @classmethod
    def read(cls, file_path: Path) -> str:
        return cls.read_result(file_path).text

    @classmethod
    def read_result(cls, file_path: Path) -> ReaderResult:
        try:
            import docx
        except ImportError:
            raise RuntimeError(
                "DOCX reader requires python-docx: pip install hexis[readers]"
            )

        doc = docx.Document(file_path)
        header = f"[Format: DOCX]\n[File: {file_path.name}]\n\n"
        parts: list[str] = []
        anchors: list[Anchor] = []
        heading_stack: list[tuple[int, str]] = []
        pos = len(header)

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style_name = str(getattr(getattr(para, "style", None), "name", "") or "")
            m = cls._HEADING_RE.match(style_name)
            if m:
                level = int(m.group(1))
                heading_stack = [h for h in heading_stack if h[0] < level]
                heading_stack.append((level, text))
                anchors.append(Anchor(
                    kind="heading",
                    value=[t for _, t in heading_stack],
                    char_offset=pos,
                ))
            parts.append(text)
            pos += len(text) + 1  # "\n" joiner

        table_count = 0
        for table in doc.tables:
            table_count += 1
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                line = "\t".join(cells)
                parts.append(line)
                pos += len(line) + 1

        warnings: list[ReaderWarning] = []
        if table_count:
            warnings.append(ReaderWarning(
                code="unsupported_feature",
                message=f"{table_count} table(s) flattened to tab-separated rows (appended after body text)",
                detail={"feature": "table_structure", "tables": table_count},
            ))

        return ReaderResult(
            text=header + "\n".join(parts),
            warnings=warnings,
            anchors=anchors,
            extractor_name="python-docx",
            extractor_version=_module_version(docx),
        )


class RtfReader(DocumentReader):
    """Reader for RTF files."""

    EXTENSIONS = {".rtf"}

    @classmethod
    def read(cls, file_path: Path) -> str:
        try:
            from striprtf.striprtf import rtf_to_text
        except ImportError:
            raise RuntimeError(
                "RTF reader requires striprtf: pip install hexis[readers]"
            )

        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            raw = f.read()
        text = rtf_to_text(raw)
        return f"[Format: RTF]\n[File: {file_path.name}]\n\n{text}"


class LatexReader(DocumentReader):
    """Reader for LaTeX .tex and .bib files (regex-based, no deps)."""

    EXTENSIONS = {".tex", ".bib"}

    @classmethod
    def read(cls, file_path: Path) -> str:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            content = f.read()

        suffix = file_path.suffix.lower()
        if suffix == ".bib":
            return cls._read_bib(content, file_path)
        return cls._read_tex(content, file_path)

    @classmethod
    def _read_tex(cls, content: str, file_path: Path) -> str:
        # Strip non-content environments
        for env in ("figure", "tikzpicture", "table", "lstlisting"):
            content = re.sub(
                rf"\\begin\{{{env}\}}.*?\\end\{{{env}\}}",
                "",
                content,
                flags=re.DOTALL,
            )
        # Strip comments
        content = re.sub(r"(?m)%.*$", "", content)
        # Strip \command{...} but keep content inside braces for text commands
        content = re.sub(r"\\(?:textbf|textit|emph|underline)\{([^}]*)\}", r"\1", content)
        # Strip remaining commands (keep braced content)
        content = re.sub(r"\\[a-zA-Z]+\*?(?:\[[^\]]*\])*\{([^}]*)\}", r"\1", content)
        # Strip standalone commands
        content = re.sub(r"\\[a-zA-Z]+\*?(?:\[[^\]]*\])*", "", content)
        # Clean up braces
        content = content.replace("{", "").replace("}", "")
        # Collapse whitespace
        content = re.sub(r"\n{3,}", "\n\n", content).strip()
        return f"[Format: LaTeX]\n[File: {file_path.name}]\n\n{content}"

    @classmethod
    def _read_bib(cls, content: str, file_path: Path) -> str:
        entries: list[str] = []
        for match in re.finditer(
            r"@\w+\{([^,]+),\s*(.*?)\n\}", content, flags=re.DOTALL
        ):
            key = match.group(1).strip()
            body = match.group(2)
            fields: dict[str, str] = {}
            for fm in re.finditer(r"(\w+)\s*=\s*\{([^}]*)\}", body):
                fields[fm.group(1).lower()] = fm.group(2).strip()
            parts = [f"[{key}]"]
            for field_name in ("author", "title", "year", "abstract"):
                if field_name in fields:
                    parts.append(f"  {field_name.title()}: {fields[field_name]}")
            entries.append("\n".join(parts))
        return (
            f"[Format: BibTeX]\n[File: {file_path.name}]\n[Entries: {len(entries)}]\n\n"
            + "\n\n".join(entries)
        )


class EmailReader(DocumentReader):
    """Reader for .eml and .mbox email files (stdlib only)."""

    EXTENSIONS = {".eml", ".mbox"}

    @classmethod
    def read(cls, file_path: Path) -> str:
        suffix = file_path.suffix.lower()
        if suffix == ".mbox":
            return cls._read_mbox(file_path)
        return cls._read_eml(file_path)

    @classmethod
    def _read_eml(cls, file_path: Path) -> str:
        import email
        import email.policy

        with open(file_path, "rb") as f:
            msg = email.message_from_bytes(f.read(), policy=email.policy.default)
        return cls._format_message(msg)

    @classmethod
    def _read_mbox(cls, file_path: Path) -> str:
        import mailbox

        mbox = mailbox.mbox(str(file_path))
        parts: list[str] = []
        for i, msg in enumerate(mbox):
            parts.append(f"--- Message {i + 1} ---")
            parts.append(cls._format_message(msg))
        if not parts:
            return f"[Email]\n[File: {file_path.name}]\n[Empty mailbox]"
        return "\n\n".join(parts)

    @classmethod
    def _format_message(cls, msg: Any) -> str:
        subject = str(msg.get("Subject", "(no subject)"))
        from_addr = str(msg.get("From", ""))
        to_addr = str(msg.get("To", ""))
        date = str(msg.get("Date", ""))

        body = ""
        if msg.is_multipart():
            for part in msg.walk():
                ct = part.get_content_type()
                if ct == "text/plain":
                    payload = part.get_payload(decode=True)
                    if payload:
                        body = payload.decode("utf-8", errors="replace")
                        break
            if not body:
                for part in msg.walk():
                    ct = part.get_content_type()
                    if ct == "text/html":
                        payload = part.get_payload(decode=True)
                        if payload:
                            body = payload.decode("utf-8", errors="replace")
                            break
        else:
            payload = msg.get_payload(decode=True)
            if payload:
                body = payload.decode("utf-8", errors="replace")
            elif isinstance(msg.get_payload(), str):
                body = msg.get_payload()

        from .textclean import clean_email_text

        header = f"[Email]\n[Subject: {subject}]\n[From: {from_addr}]\n[To: {to_addr}]\n[Date: {date}]"
        # Same cleanup as imported Gmail: markup, entities, padding, tracking links.
        return f"{header}\n\n{clean_email_text(body)}"


class EpubReader(DocumentReader):
    """Reader for EPUB e-book files."""

    EXTENSIONS = {".epub"}

    @classmethod
    def read(cls, file_path: Path) -> str:
        try:
            import ebooklib
            from ebooklib import epub
        except ImportError:
            raise RuntimeError(
                "EPUB reader requires ebooklib and beautifulsoup4: pip install hexis[readers]"
            )

        try:
            from bs4 import BeautifulSoup
        except ImportError:
            raise RuntimeError(
                "EPUB reader requires beautifulsoup4: pip install hexis[readers]"
            )

        book = epub.read_epub(str(file_path), options={"ignore_ncx": True})
        title = book.get_metadata("DC", "title")
        title_str = title[0][0] if title else file_path.stem
        author = book.get_metadata("DC", "creator")
        author_str = author[0][0] if author else "Unknown"

        parts: list[str] = []
        for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
            soup = BeautifulSoup(item.get_body_content(), "html.parser")
            text = soup.get_text(separator="\n", strip=True)
            if text.strip():
                parts.append(text)

        header = f"[Format: EPUB]\n[Title: {title_str}]\n[Author: {author_str}]"
        return f"{header}\n\n" + "\n\n".join(parts)


class PptxReader(DocumentReader):
    """Reader for PowerPoint .pptx files."""

    EXTENSIONS = {".pptx"}

    @classmethod
    def read(cls, file_path: Path) -> str:
        try:
            from pptx import Presentation
        except ImportError:
            raise RuntimeError(
                "PPTX reader requires python-pptx: pip install hexis[readers]"
            )

        prs = Presentation(str(file_path))
        parts: list[str] = []
        for i, slide in enumerate(prs.slides):
            slide_texts: list[str] = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if text:
                            slide_texts.append(text)
            if slide_texts:
                parts.append(f"[Slide {i + 1}]\n" + "\n".join(slide_texts))

        header = f"[Format: PPTX]\n[File: {file_path.name}]\n[Slides: {len(prs.slides)}]"
        return f"{header}\n\n" + "\n\n".join(parts)


class XlsxReader(DocumentReader):
    """Reader for Excel .xlsx files. Legacy .xls is not supported (openpyxl
    cannot read it) — get_reader fails loud with a convert-to-xlsx hint."""

    EXTENSIONS = {".xlsx"}

    # Overridden from config ingest.xlsx_max_rows_per_sheet by the pipeline.
    # Capping is never silent: it always emits a truncated_rows warning.
    MAX_ROWS_PER_SHEET = 5000

    @classmethod
    def read(cls, file_path: Path) -> str:
        return cls.read_result(file_path).text

    @classmethod
    def read_result(cls, file_path: Path) -> ReaderResult:
        try:
            import openpyxl
        except ImportError:
            raise RuntimeError(
                "XLSX reader requires openpyxl: pip install hexis[readers]"
            )

        wb = openpyxl.load_workbook(str(file_path), read_only=True, data_only=True)
        sheet_names = list(wb.sheetnames)
        header = f"[Format: XLSX]\n[File: {file_path.name}]\n[Sheets: {len(sheet_names)}]"
        parts: list[str] = []
        anchors: list[Anchor] = []
        warnings: list[ReaderWarning] = []
        pos = len(header) + 2  # "\n\n" after the header

        for sheet_name in sheet_names:
            ws = wb[sheet_name]
            if parts:
                pos += 2  # "\n\n" between sheets
            sheet_marker = f"[Sheet: {sheet_name}]"
            anchors.append(Anchor(kind="sheet", value=sheet_name, char_offset=pos))
            lines: list[str] = [sheet_marker]
            pos += len(sheet_marker) + 1  # marker + "\n"
            emitted = 0
            total = 0
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                total += 1
                if i >= cls.MAX_ROWS_PER_SHEET:
                    continue  # keep counting for the warning total
                cells = [str(c) if c is not None else "" for c in row]
                line = "\t".join(cells)
                anchors.append(Anchor(kind="row", value=i + 1, char_offset=pos))
                lines.append(line)
                pos += len(line) + 1
                emitted += 1
            if total > emitted:
                notice = f"... (truncated: showing {emitted} of {total} rows)"
                lines.append(notice)
                pos += len(notice) + 1
                warnings.append(ReaderWarning(
                    code="truncated_rows",
                    message=f"Sheet '{sheet_name}' truncated to {emitted} of {total} rows "
                            f"(raise ingest.xlsx_max_rows_per_sheet to keep more)",
                    detail={"sheet": sheet_name, "emitted": emitted, "total": total},
                ))
            pos -= 1  # the last line has no trailing "\n" inside the join
            parts.append("\n".join(lines))
        wb.close()

        return ReaderResult(
            text=f"{header}\n\n" + "\n\n".join(parts),
            warnings=warnings,
            anchors=anchors,
            extractor_name="openpyxl",
            extractor_version=_module_version(openpyxl),
            # Formulas are read as their last-calculated values (data_only).
            metadata={"formulas": "values_only", "sheets": sheet_names},
        )


class NotebookReader(DocumentReader):
    """Reader for Jupyter .ipynb notebooks (stdlib json only)."""

    EXTENSIONS = {".ipynb"}

    @classmethod
    def read(cls, file_path: Path) -> str:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            nb = json.loads(f.read())

        cells = nb.get("cells", [])
        kernel = (
            nb.get("metadata", {}).get("kernelspec", {}).get("language", "python")
        )

        parts: list[str] = []
        for cell in cells:
            cell_type = cell.get("cell_type", "")
            source = "".join(cell.get("source", []))
            if not source.strip():
                continue
            if cell_type == "markdown":
                parts.append(source)
            elif cell_type == "code":
                parts.append(f"```{kernel}\n{source}\n```")

        header = f"[Format: Jupyter Notebook]\n[File: {file_path.name}]\n[Cells: {len(cells)}]"
        return f"{header}\n\n" + "\n\n".join(parts)


class RssReader(DocumentReader):
    """Reader for RSS/Atom feeds (URL-based, not file-based)."""

    @classmethod
    def read(cls, url: str) -> str:  # type: ignore[override]
        try:
            import feedparser
        except ImportError:
            raise RuntimeError(
                "RSS reader requires feedparser: pip install hexis[readers]"
            )

        feed = feedparser.parse(url)
        if not feed.entries:
            return ""

        feed_title = feed.feed.get("title", url)
        parts: list[str] = [f"[Format: RSS/Atom]\n[Feed: {feed_title}]\n[Entries: {len(feed.entries)}]"]
        for entry in feed.entries:
            title = entry.get("title", "(no title)")
            link = entry.get("link", "")
            summary = entry.get("summary", "")
            published = entry.get("published", "")
            # Strip HTML from summary
            summary = re.sub(r"<[^>]+>", "", summary)
            entry_parts = [f"## {title}"]
            if published:
                entry_parts.append(f"Published: {published}")
            if link:
                entry_parts.append(f"Link: {link}")
            if summary:
                entry_parts.append(summary)
            parts.append("\n".join(entry_parts))

        return "\n\n".join(parts)


class ZipArchiveReader(DocumentReader):
    """Reader for .zip bundles of notes, documents and exported mail.

    Reads the text-like entries (plain text, Markdown, CSV, JSON, .eml) into
    one document, one section per entry, and skips everything else. A guard
    stops at 50 MB uncompressed so a large archive can't exhaust memory.
    """

    EXTENSIONS = {".zip"}
    MAX_TOTAL_BYTES = 50 * 1024 * 1024
    TEXT_SUFFIXES = {".txt", ".text", ".md", ".markdown", ".csv", ".json", ".eml"}

    @classmethod
    def read(cls, file_path: Path) -> str:
        import email
        import email.policy
        import zipfile

        if not zipfile.is_zipfile(file_path):
            raise RuntimeError(f"Invalid .zip archive: {file_path.name}")

        sections: list[str] = []
        skipped = 0
        total_bytes = 0
        with zipfile.ZipFile(file_path, "r") as zf:
            for info in sorted(zf.infolist(), key=lambda i: i.filename):
                if info.is_dir():
                    continue
                entry = Path(info.filename)
                if entry.name.startswith(".") or "__MACOSX" in entry.parts:
                    continue
                suffix = entry.suffix.lower()
                if suffix not in cls.TEXT_SUFFIXES:
                    skipped += 1
                    continue
                if total_bytes + info.file_size > cls.MAX_TOTAL_BYTES:
                    sections.append("[Archive truncated: reached the 50 MB uncompressed limit]")
                    break
                raw = zf.read(info.filename)
                total_bytes += len(raw)
                if suffix == ".eml":
                    msg = email.message_from_bytes(raw, policy=email.policy.default)
                    text = EmailReader._format_message(msg)
                else:
                    text = raw.decode("utf-8", errors="replace")
                if text.strip():
                    sections.append(f"--- Entry: {info.filename} ---\n{text.strip()}")

        header = f"[Format: ZIP Archive]\n[File: {file_path.name}]"
        if skipped:
            header += f"\n[Skipped {skipped} entries that are not text, Markdown, CSV, JSON or email]"
        if not sections:
            return f"{header}\n[No readable entries found]"
        return f"{header}\n\n" + "\n\n".join(sections)


def get_reader(file_path: Path) -> DocumentReader:
    suffix = file_path.suffix.lower()
    if suffix == ".xls":
        # openpyxl cannot read legacy .xls; claiming support would just crash
        # inside the parser. Fail loud with the fix.
        raise RuntimeError(
            "Legacy .xls is not supported; convert it to .xlsx "
            "(e.g. with LibreOffice: soffice --convert-to xlsx) and re-ingest."
        )
    if suffix == ".pdf":
        return PDFReader()
    if suffix in [".md", ".markdown"]:
        return MarkdownReader()
    if suffix in DocxReader.EXTENSIONS:
        return DocxReader()
    if suffix in RtfReader.EXTENSIONS:
        return RtfReader()
    if suffix in LatexReader.EXTENSIONS:
        return LatexReader()
    if suffix in EmailReader.EXTENSIONS:
        return EmailReader()
    if suffix in EpubReader.EXTENSIONS:
        return EpubReader()
    if suffix in PptxReader.EXTENSIONS:
        return PptxReader()
    if suffix in XlsxReader.EXTENSIONS:
        return XlsxReader()
    if suffix in NotebookReader.EXTENSIONS:
        return NotebookReader()
    if suffix in ZipArchiveReader.EXTENSIONS:
        return ZipArchiveReader()
    if suffix in DataReader.DATA_EXTENSIONS:
        return DataReader()
    if suffix in ImageReader.IMAGE_EXTENSIONS:
        return ImageReader()
    if suffix in AudioReader.AUDIO_EXTENSIONS:
        return AudioReader()
    if suffix in VideoReader.VIDEO_EXTENSIONS:
        return VideoReader()
    if suffix in CodeReader.LANGUAGE_MAP:
        return CodeReader()
    return TextReader()
